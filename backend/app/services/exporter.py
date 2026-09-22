"""成绩导出：按评分细则（项目 / 作业）导出 Excel，可选限定某个班级。"""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from ..models import Class, CourseClass, GradingResult, ResultStatus, Rubric, Student

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

STATUS_LABEL = {
    ResultStatus.pending: "待评审",
    ResultStatus.grading: "评审中",
    ResultStatus.graded: "已评审",
    ResultStatus.failed: "评审失败",
    "missing": "未提交",
}

KIND_LABEL = {"homework": "作业", "lab_report": "实验报告"}


def _style_header(sheet, columns: int, row: int = 1) -> None:
    for col in range(1, columns + 1):
        cell = sheet.cell(row=row, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER


def _autosize(sheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def load_rows(
    db: Session,
    rubric_id: int,
    *,
    class_id: int | None = None,
    course_id: int | None = None,
    include_missing: bool = True,
) -> list[dict]:
    """把评分结果整理成导出行；include_missing 时补上没交的学生。"""
    stmt = (
        select(GradingResult)
        .options(
            selectinload(GradingResult.student).selectinload(Student.class_),
            selectinload(GradingResult.class_),
        )
        .where(GradingResult.rubric_id == rubric_id)
    )
    if class_id is not None:
        stmt = stmt.where(GradingResult.class_id == class_id)
    if course_id is not None:
        stmt = stmt.join(
            CourseClass, CourseClass.class_id == GradingResult.class_id
        ).where(CourseClass.course_id == course_id)
    results = db.execute(stmt).scalars().all()

    rows: list[dict] = []
    for item in results:
        klass = item.class_
        rows.append(
            {
                "department": klass.department if klass else "",
                "major": klass.major if klass else "",
                "class_name": klass.name if klass else "",
                "class_id": item.class_id,
                "student_no": item.student.student_no if item.student else "",
                "student_name": item.student.name if item.student else "",
                "filename": item.source_filename,
                "status": STATUS_LABEL.get(item.status, item.status.value),
                "score": item.score,
                "comment": item.comment or "",
                "graded_at": item.graded_at.strftime("%Y-%m-%d %H:%M") if item.graded_at else "",
                "model": item.model or "",
                "is_manual": item.is_manual,
                "sort_key": (klass.full_name if klass else "", item.student.student_no if item.student else ""),
            }
        )

    if include_missing:
        involved: set[int] = {item["class_id"] for item in rows}
        if class_id is not None:
            involved.add(class_id)
        if course_id is not None:
            involved.update(
                db.execute(
                    select(CourseClass.class_id).where(CourseClass.course_id == course_id)
                ).scalars().all()
            )
        submitted = set(
            db.execute(
                select(GradingResult.student_id).where(GradingResult.rubric_id == rubric_id)
            ).scalars().all()
        )

        for cid in sorted(involved):
            klass = db.get(Class, cid)
            if klass is None:
                continue
            students = db.execute(
                select(Student).where(Student.class_id == cid).order_by(Student.student_no)
            ).scalars().all()
            for student in students:
                if student.id in submitted:
                    continue
                rows.append(
                    {
                        "department": klass.department,
                        "major": klass.major,
                        "class_name": klass.name,
                        "class_id": cid,
                        "student_no": student.student_no,
                        "student_name": student.name,
                        "filename": "",
                        "status": STATUS_LABEL["missing"],
                        "score": None,
                        "comment": "",
                        "graded_at": "",
                        "model": "",
                        "is_manual": False,
                        "sort_key": (klass.full_name, student.student_no),
                    }
                )

    rows.sort(key=lambda item: item["sort_key"])
    return rows


def build_workbook(db: Session, rubric: Rubric, rows: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "成绩与评语"

    headers = ["序号", "院系", "专业", "班级", "学号", "姓名", "提交文件", "状态", "总分", "评语", "评审时间"]
    sheet.append(headers)
    _style_header(sheet, len(headers))

    for index, row in enumerate(rows, start=1):
        sheet.append(
            [
                index,
                row["department"],
                row["major"],
                row["class_name"],
                row["student_no"],
                row["student_name"],
                row["filename"],
                row["status"],
                row["score"],
                row["comment"],
                row["graded_at"],
            ]
        )

    _autosize(sheet, [6, 16, 18, 14, 16, 12, 34, 10, 8, 60, 18])
    sheet.freeze_panes = "A2"
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (7, 10))

    # ---- 班级统计 ----
    stats_sheet = workbook.create_sheet("班级统计")
    stats_sheet.append(["班级", "应交人数", "已评审", "未提交", "平均分", "最高分", "最低分"])
    _style_header(stats_sheet, 7)

    buckets: dict[str, list[dict]] = {}
    for row in rows:
        key = f"{row['department']}-{row['major']}-{row['class_name']}".strip("-")
        buckets.setdefault(key, []).append(row)

    for key in sorted(buckets):
        group = buckets[key]
        scores = [item["score"] for item in group if item["score"] is not None]
        graded = sum(1 for item in group if item["status"] == STATUS_LABEL[ResultStatus.graded])
        missing = sum(1 for item in group if item["status"] == STATUS_LABEL["missing"])
        stats_sheet.append(
            [
                key,
                len(group),
                graded,
                missing,
                round(sum(scores) / len(scores), 2) if scores else "",
                max(scores) if scores else "",
                min(scores) if scores else "",
            ]
        )
    _autosize(stats_sheet, [34, 12, 10, 10, 10, 10, 10])
    for row in stats_sheet.iter_rows(min_row=1, max_row=stats_sheet.max_row, max_col=7):
        for cell in row:
            cell.border = _BORDER

    # ---- 概览 ----
    info_sheet = workbook.create_sheet("概览")
    scores = [row["score"] for row in rows if row["score"] is not None]
    info_rows = [
        ("项目 / 作业", rubric.name),
        ("类型", KIND_LABEL.get(rubric.kind.value, rubric.kind.value)),
        ("满分", rubric.total_score),
        ("记录数", len(rows)),
        ("已评审", sum(1 for row in rows if row["status"] == STATUS_LABEL[ResultStatus.graded])),
        ("未提交", sum(1 for row in rows if row["status"] == STATUS_LABEL["missing"])),
        ("平均分", round(sum(scores) / len(scores), 2) if scores else ""),
        ("最高分", max(scores) if scores else ""),
        ("最低分", min(scores) if scores else ""),
        ("导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    info_sheet.append(["项目", "值"])
    _style_header(info_sheet, 2)
    for entry in info_rows:
        info_sheet.append(list(entry))
    _autosize(info_sheet, [16, 46])
    for row in info_sheet.iter_rows(min_row=1, max_row=info_sheet.max_row, max_col=2):
        for cell in row:
            cell.border = _BORDER

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_filename(rubric: Rubric, class_name: str | None = None) -> str:
    stamp = datetime.now().strftime("%Y%m%d%H%M")
    safe_name = "".join(ch for ch in rubric.name if ch not in '\\/:*?"<>|').strip() or "成绩"
    scope = f"_{class_name}" if class_name else ""
    return f"{safe_name}{scope}_成绩评语_{stamp}.xlsx"
