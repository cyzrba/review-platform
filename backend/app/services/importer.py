"""班级 / 学生名单导入（CSV / Excel）。

名单文件约定列：学号(工号)、姓名、院系、专业、班级、加入时间、入学年份。
学号和姓名必填，其余可为空。
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

# 表头别名 -> 标准字段
ALIASES: dict[str, str] = {
    # 学号 / 工号
    "学号": "student_no", "工号": "student_no", "编号": "student_no", "学籍号": "student_no",
    "学生学号": "student_no", "no": "student_no", "id": "student_no", "student_no": "student_no",
    "studentno": "student_no", "number": "student_no",
    # 姓名
    "姓名": "name", "名字": "name", "学生姓名": "name", "学生名称": "name", "name": "name",
    "student_name": "name", "username": "name",
    # 院系
    "院系": "department", "学院": "department", "系": "department", "系别": "department",
    "department": "department", "faculty": "department", "college": "department", "school": "department",
    # 专业
    "专业": "major", "专业名称": "major", "major": "major", "specialty": "major",
    # 班级
    "班级": "class_name", "班级名称": "class_name", "班号": "class_name", "教学班": "class_name",
    "行政班": "class_name", "class": "class_name", "class_name": "class_name", "classname": "class_name",
    # 加入时间
    "加入时间": "joined_at", "入班时间": "joined_at", "加入日期": "joined_at", "入班日期": "joined_at",
    "进班时间": "joined_at", "join_date": "joined_at", "joined_at": "joined_at", "created": "joined_at",
    # 入学年份
    "入学年份": "enrollment_year", "入学年": "enrollment_year", "年级": "enrollment_year",
    "入学年度": "enrollment_year", "enrollment_year": "enrollment_year", "grade": "enrollment_year",
    "year": "enrollment_year",
    # 邮箱
    "邮箱": "email", "电子邮箱": "email", "email": "email", "mail": "email",
}

REQUIRED_FIELDS = ("student_no", "name")

_DATE_FORMATS = (
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%Y.%m.%d",
    "%Y/%m/%d %H:%M:%S",
    "%Y年%m月%d日",
    "%Y-%m-%dT%H:%M:%S",
    "%Y%m%d",
)
_YEAR_RE = re.compile(r"(19\d{2}|20\d{2}|21\d{2})")


class ImportError_(ValueError):
    """导入文件无法解析。"""


@dataclass
class ParsedRoster:
    rows: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------- #
# 单元格值归一化
# --------------------------------------------------------------------------- #
def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _student_no(value: Any) -> str:
    text = _text(value)
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text


def parse_joined_at(value: Any) -> datetime | None:
    """支持 Excel 日期单元格和常见字符串格式。"""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    text = _text(value)
    if not text:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    # 退一步：从字符串里抠出年月日
    parts = re.findall(r"\d+", text)
    if len(parts) >= 3:
        try:
            return datetime(int(parts[0]), int(parts[1]), int(parts[2]))
        except ValueError:
            return None
    return None


def parse_enrollment_year(value: Any) -> int | None:
    """'2023'、'2023级'、'2023年' 都能解析。"""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.year
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        year = int(value)
        return year if 1900 <= year <= 2200 else None
    match = _YEAR_RE.search(_text(value))
    return int(match.group(1)) if match else None


# --------------------------------------------------------------------------- #
# 文件解析
# --------------------------------------------------------------------------- #
def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gbk", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _parse_csv(data: bytes) -> list[dict[str, Any]]:
    text = _decode_text(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [dict(row) for row in csv.DictReader(io.StringIO(text), dialect=dialect)]


def _parse_xlsx(data: bytes) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    if sheet is None:
        return []
    rows = sheet.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        workbook.close()
        return []
    columns = [_text(cell) for cell in header]

    out: list[dict[str, Any]] = []
    for raw in rows:
        record: dict[str, Any] = {}
        for index, value in enumerate(raw):
            if index >= len(columns):
                break
            record[columns[index]] = value
        if any(_text(value) for value in record.values()):
            out.append(record)
    workbook.close()
    return out


def parse_roster(filename: str, data: bytes) -> ParsedRoster:
    """把名单文件解析成统一结构。"""
    lower = (filename or "").lower()
    if lower.endswith((".xlsx", ".xlsm")):
        raw_rows = _parse_xlsx(data)
    elif lower.endswith((".csv", ".txt")):
        raw_rows = _parse_csv(data)
    elif lower.endswith(".xls"):
        raise ImportError_("暂不支持 .xls，请另存为 .xlsx 或 .csv")
    else:
        raise ImportError_("只支持 .csv / .xlsx 名单文件")

    if not raw_rows:
        raise ImportError_("名单文件里没有数据行")

    header_map: dict[str, str] = {}
    for raw_header in raw_rows[0]:
        key = _text(raw_header).lower().replace(" ", "")
        target = ALIASES.get(key)
        if target and target not in header_map.values():
            header_map[raw_header] = target

    missing = [field for field in REQUIRED_FIELDS if field not in header_map.values()]
    if missing:
        raise ImportError_(
            "表头必须包含「学号/工号」和「姓名」两列（可选：院系、专业、班级、加入时间、入学年份）"
        )

    parsed = ParsedRoster()
    seen: set[str] = set()
    for index, row in enumerate(raw_rows, start=2):
        record: dict[str, Any] = {field: None for field in set(ALIASES.values())}
        for source, target in header_map.items():
            record[target] = row.get(source)

        student_no = _student_no(record.get("student_no"))
        name = _text(record.get("name"))
        if not student_no or not name:
            parsed.warnings.append(f"第 {index} 行缺少学号或姓名，已跳过")
            continue
        if student_no in seen:
            parsed.warnings.append(f"第 {index} 行学号 {student_no} 在文件内重复，已跳过")
            continue
        seen.add(student_no)

        parsed.rows.append(
            {
                "student_no": student_no,
                "name": name,
                "department": _text(record.get("department")),
                "major": _text(record.get("major")),
                "class_name": _text(record.get("class_name")),
                "joined_at": parse_joined_at(record.get("joined_at")),
                "enrollment_year": parse_enrollment_year(record.get("enrollment_year")),
                "email": _text(record.get("email")) or None,
            }
        )

    if not parsed.rows:
        raise ImportError_("名单文件里没有有效数据行")
    return parsed


ROSTER_TEMPLATE_CSV = (
    "学号,姓名,院系,专业,班级,加入时间,入学年份\n"
    "2023010101,张三,计算机学院,软件工程,软工2301,2023-09-01,2023\n"
    "2023010102,李四,计算机学院,软件工程,软工2301,2023-09-01,2023\n"
    "2023010103,王五,计算机学院,计算机科学与技术,计科2302,2023-09-01,2023\n"
)
