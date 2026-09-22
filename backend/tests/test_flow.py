"""端到端：建课程 → 导入班级与学生 → 建评分细则 → 上传 → AI 评审 → 导出。

每个用例用独立的课程名、院系名和学号，避免用例之间互相干扰。
"""

from __future__ import annotations

import io
import itertools
import zipfile
from dataclasses import dataclass

_counter = itertools.count(1)


@dataclass
class Fixture:
    tag: str
    course_id: int
    department: str
    class_soft: int  # 软工2301（张三、李四）
    class_cs: int  # 计科2302（王五）
    no_zhang: str
    no_li: str
    no_wang: str


def items(response) -> list:
    """列表接口统一是分页信封，这里取 items。"""
    body = response.json()
    assert "items" in body, body
    return body["items"]


def setup(client) -> Fixture:
    index = next(_counter)
    tag = str(index)
    department = f"计算机学院{tag}"
    no_zhang = f"2023{index:04d}01"
    no_li = f"2023{index:04d}02"
    no_wang = f"2023{index:04d}03"

    course = client.post(
        "/api/courses",
        json={"name": f"数据结构-{tag}", "code": f"CS{index:03d}", "term": "2024-2025-1"},
    )
    assert course.status_code == 201, course.text
    course_id = course.json()["id"]

    roster = (
        "学号,姓名,院系,专业,班级,加入时间,入学年份\n"
        f"{no_zhang},张三,{department},软件工程,软工2301,2023-09-01,2023\n"
        f"{no_li},李四,{department},软件工程,软工2301,2023-09-01,2023\n"
        f"{no_wang},王五,{department},计算机科学与技术,计科2302,2023-09-01,2023\n"
    ).encode("utf-8-sig")

    imported = client.post(
        "/api/classes/import",
        files={"file": ("roster.csv", roster, "text/csv")},
        data={"course_id": str(course_id)},
    )
    assert imported.status_code == 200, imported.text
    body = imported.json()
    assert body["classes_created"] == 2
    assert body["students_created"] == 3

    classes = items(client.get("/api/courses/{}/classes".format(course_id), params={"page_size": 50}))
    assert len(classes) == 2
    by_name = {item["name"]: item["id"] for item in classes}
    return Fixture(
        tag=tag,
        course_id=course_id,
        department=department,
        class_soft=by_name["软工2301"],
        class_cs=by_name["计科2302"],
        no_zhang=no_zhang,
        no_li=no_li,
        no_wang=no_wang,
    )


def zip_bytes(entries: dict) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in entries.items():
            archive.writestr(name, payload)
    return buffer.getvalue()


def create_rubric(
    client,
    *,
    kind: str,
    name: str,
    course_ids: list[int] | None = None,
    with_criteria: bool = True,
    with_description: bool = True,
) -> int:
    response = client.post(
        "/api/rubrics",
        json={
            "name": name,
            "kind": kind,
            "description": "完成实验并撰写报告" if with_description else None,
            "criteria": "按完成度给分：内容完整 60 分，结论合理 40 分。" if with_criteria else None,
            "total_score": 100,
            "course_ids": course_ids or [],
        },
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def upload(client, rubric_id: int, filename: str, payload: bytes, **extra):
    data = {key: str(value) for key, value in extra.items() if value is not None} or None
    response = client.post(
        f"/api/rubrics/{rubric_id}/submissions",
        files={"files": (filename, payload, "application/zip")},
        data=data,
    )
    assert response.status_code == 200, response.text
    return response.json()


def grade_all(client, rubric_id: int, **payload) -> dict:
    response = client.post(
        f"/api/rubrics/{rubric_id}/grade", json=payload, params={"sync": "true"}
    )
    assert response.status_code == 200, response.text
    return response.json()


def results_of(client, rubric_id: int, **params) -> list:
    params.setdefault("page_size", 50)
    return items(client.get(f"/api/rubrics/{rubric_id}/results", params=params))


# --------------------------------------------------------------------------- #
# 课程与班级挂靠
# --------------------------------------------------------------------------- #
def test_course_crud_and_linking(client):
    fixture = setup(client)

    course = client.get(f"/api/courses/{fixture.course_id}").json()
    assert course["class_count"] == 2
    assert course["student_count"] == 3
    assert course["rubric_count"] == 0

    # 同一个班级可以挂到另一个课程下（多对多）
    other = client.post("/api/courses", json={"name": f"操作系统-{fixture.tag}"}).json()
    linked = client.post(
        f"/api/courses/{other['id']}/classes", json={"class_ids": [fixture.class_soft]}
    ).json()
    assert linked["linked"] == 1

    klass = client.get(f"/api/classes/{fixture.class_soft}").json()
    assert {item["name"] for item in klass["courses"]} == {
        f"数据结构-{fixture.tag}",
        f"操作系统-{fixture.tag}",
    }

    # 重复挂载会跳过
    again = client.post(
        f"/api/courses/{other['id']}/classes", json={"class_ids": [fixture.class_soft]}
    ).json()
    assert again["linked"] == 0
    assert again["skipped"] == 1

    # 摘掉关联不影响班级本身
    assert client.delete(f"/api/courses/{other['id']}/classes/{fixture.class_soft}").status_code == 204
    assert client.get(f"/api/classes/{fixture.class_soft}").status_code == 200

    # 删除课程不会删掉班级
    assert client.delete(f"/api/courses/{other['id']}").status_code == 204
    assert client.get(f"/api/classes/{fixture.class_soft}").status_code == 200


def test_course_rubric_linking(client):
    fixture = setup(client)
    rubric_id = create_rubric(
        client, kind="lab_report", name=f"实验一-{fixture.tag}", course_ids=[fixture.course_id]
    )

    rubric = client.get(f"/api/rubrics/{rubric_id}").json()
    assert [item["id"] for item in rubric["courses"]] == [fixture.course_id]

    listed = items(
        client.get("/api/courses/{}/rubrics".format(fixture.course_id), params={"page_size": 50})
    )
    assert [item["id"] for item in listed] == [rubric_id]

    course = client.get(f"/api/courses/{fixture.course_id}").json()
    assert course["rubric_count"] == 1

    # 改挂到别的课程
    other = client.post("/api/courses", json={"name": f"编译原理-{fixture.tag}"}).json()
    updated = client.patch(f"/api/rubrics/{rubric_id}", json={"course_ids": [other["id"]]}).json()
    assert [item["id"] for item in updated["courses"]] == [other["id"]]

    # 摘掉关联不影响细则本身
    assert client.delete(f"/api/courses/{other['id']}/rubrics/{rubric_id}").status_code == 204
    assert client.get(f"/api/rubrics/{rubric_id}").json()["courses"] == []


def test_import_requires_course(client):
    roster = "学号,姓名,班级\n2023010101,张三,软工2301\n".encode("utf-8-sig")
    response = client.post(
        "/api/classes/import", files={"file": ("roster.csv", roster, "text/csv")}
    )
    assert response.status_code == 422  # 缺少必填的 course_id

    response = client.post(
        "/api/classes/import",
        files={"file": ("roster.csv", roster, "text/csv")},
        data={"course_id": "999999"},
    )
    assert response.status_code == 404
    assert "课程不存在" in response.json()["detail"]


def test_student_and_class_fields(client):
    fixture = setup(client)
    students = items(
        client.get(
            f"/api/classes/{fixture.class_soft}/students", params={"page_size": 50}
        )
    )
    assert len(students) == 2
    first = students[0]
    assert first["joined_at"].startswith("2023-09-01")
    assert first["enrollment_year"] == 2023
    assert first["department"] == fixture.department
    assert first["major"] == "软件工程"
    assert first["class_name"] == "软工2301"

    klass = client.get(f"/api/classes/{fixture.class_soft}").json()
    assert klass["full_name"] == f"{fixture.department}-软件工程-软工2301"


# --------------------------------------------------------------------------- #
# 分页
# --------------------------------------------------------------------------- #
def test_all_list_endpoints_are_paginated(client):
    fixture = setup(client)
    rubric_id = create_rubric(client, kind="lab_report", name=f"实验分页-{fixture.tag}")

    first_page = client.get("/api/courses", params={"page": 1, "page_size": 1}).json()
    assert set(first_page) == {"items", "total", "page", "page_size", "pages"}
    assert len(first_page["items"]) == 1
    assert first_page["total"] >= 2  # 本用例 + 其它用例建的课程
    assert first_page["pages"] == (first_page["total"] + 0) // 1

    second_page = client.get("/api/courses", params={"page": 2, "page_size": 1}).json()
    assert second_page["items"][0]["id"] != first_page["items"][0]["id"]

    for url, params in (
        ("/api/classes", {"page_size": 2}),
        ("/api/students", {"page_size": 2}),
        ("/api/rubrics", {"page_size": 2}),
        (f"/api/courses/{fixture.course_id}/classes", {"page_size": 1}),
        (f"/api/rubrics/{rubric_id}/results", {"page_size": 1}),
        (f"/api/rubrics/{rubric_id}/classes", {"page_size": 1}),
        (f"/api/rubrics/{rubric_id}/class-summary", {"page_size": 1}),
    ):
        body = client.get(url, params=params).json()
        assert set(body) == {"items", "total", "page", "page_size", "pages"}, url
        assert body["page_size"] == params["page_size"], url
        assert len(body["items"]) <= params["page_size"], url

    # page_size 超上限要被拒绝
    assert client.get("/api/classes", params={"page_size": 9999}).status_code == 422
    assert client.get("/api/classes", params={"page": 0}).status_code == 422


def test_results_pagination_walks_all_pages(client):
    fixture = setup(client)
    rubric_id = create_rubric(client, kind="lab_report", name=f"实验翻页-{fixture.tag}")
    upload(
        client,
        rubric_id,
        "批次.zip",
        zip_bytes(
            {
                f"张三_{fixture.no_zhang}.zip": zip_bytes(
                    {f"张三_{fixture.no_zhang}.pdf": b"%PDF-1.4"}
                ),
                f"李四_{fixture.no_li}.zip": zip_bytes(
                    {f"李四_{fixture.no_li}.pdf": b"%PDF-1.4"}
                ),
            }
        ),
    )

    page1 = client.get(f"/api/rubrics/{rubric_id}/results", params={"page": 1, "page_size": 1}).json()
    assert page1["total"] == 2
    assert page1["pages"] == 2
    assert len(page1["items"]) == 1

    page2 = client.get(f"/api/rubrics/{rubric_id}/results", params={"page": 2, "page_size": 1}).json()
    assert len(page2["items"]) == 1
    assert page1["items"][0]["id"] != page2["items"][0]["id"]

    empty = client.get(f"/api/rubrics/{rubric_id}/results", params={"page": 9, "page_size": 1}).json()
    assert empty["items"] == []
    assert empty["total"] == 2


# --------------------------------------------------------------------------- #
# 实验报告：学生级压缩包 → pdf
# --------------------------------------------------------------------------- #
def test_lab_report_pipeline(client):
    fixture = setup(client)
    rubric_id = create_rubric(
        client, kind="lab_report", name=f"实验一-{fixture.tag}", course_ids=[fixture.course_id]
    )

    hint = client.get(f"/api/rubrics/{rubric_id}/upload-hint").json()
    assert hint["kind"] == "lab_report"
    assert "学生名字_学号.zip" in hint["expect"]

    outer = zip_bytes(
        {
            f"张三_{fixture.no_zhang}.zip": zip_bytes(
                {f"张三_{fixture.no_zhang}.pdf": b"%PDF-1.4 fake"}
            ),
            f"李四_{fixture.no_li}.zip": zip_bytes(
                {f"李四_{fixture.no_li}.pdf": b"%PDF-1.4 fake"}
            ),
            "无名氏.zip": zip_bytes({"报告.pdf": b"%PDF-1.4 fake"}),
        }
    )
    summary = upload(client, rubric_id, "实验一批次.zip", outer, course_id=fixture.course_id)
    assert summary["scanned_files"] == 3
    assert summary["accepted_files"] == 3
    assert summary["matched"] == 2
    assert summary["created"] == 2
    assert any("无名氏" in issue["filename"] for issue in summary["issues"])

    rows = results_of(client, rubric_id)
    assert len(rows) == 2
    assert {item["student_no"] for item in rows} == {fixture.no_zhang, fixture.no_li}
    assert all(item["status"] == "pending" for item in rows)
    assert all(item["source_filename"].endswith(".pdf") for item in rows)
    assert all(item["class_id"] == fixture.class_soft for item in rows)

    assert grade_all(client, rubric_id)["queued"] == 2

    rows = results_of(client, rubric_id)
    assert all(item["status"] == "graded" for item in rows)
    assert all(0 < item["score"] <= 100 for item in rows)
    assert all(item["comment"] for item in rows)
    assert all(item["graded_at"] for item in rows)

    # 按课程过滤
    assert len(results_of(client, rubric_id, course_id=fixture.course_id)) == 2

    summaries = items(client.get(f"/api/rubrics/{rubric_id}/class-summary"))
    assert len(summaries) == 1
    assert summaries[0]["graded"] == 2
    assert summaries[0]["average_score"] is not None

    response = client.get(f"/api/rubrics/{rubric_id}/export.xlsx")
    assert response.status_code == 200
    assert response.content[:2] == b"PK"

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook["成绩与评语"]
    headers = [cell.value for cell in sheet[1]]
    assert headers[4:6] == ["学号", "姓名"]
    assert sheet.max_row == 3  # 表头 + 张三 + 李四
    assert "班级统计" in workbook.sheetnames

    csv_response = client.get(f"/api/rubrics/{rubric_id}/export.csv", params={"course_id": fixture.course_id})
    assert "张三" in csv_response.content.decode("utf-8-sig")


# --------------------------------------------------------------------------- #
# 作业：压缩包里的 word 文件
# --------------------------------------------------------------------------- #
def test_homework_pipeline(client):
    fixture = setup(client)
    rubric_id = create_rubric(client, kind="homework", name=f"作业一-{fixture.tag}")

    hint = client.get(f"/api/rubrics/{rubric_id}/upload-hint").json()
    assert "院系-专业-班级名称-学号-学生名称" in hint["expect"]

    outer = zip_bytes(
        {
            f"{fixture.department}-软件工程-软工2301-{fixture.no_zhang}-张三.docx": b"docx",
            "汇总包.zip": zip_bytes(
                {
                    f"{fixture.department}-计算机科学与技术-计科2302-{fixture.no_wang}-王五.docx": b"docx"
                }
            ),
            "无关说明.txt": "忽略",
        }
    )
    summary = upload(client, rubric_id, "作业一.zip", outer)
    assert summary["scanned_files"] == 3
    assert summary["accepted_files"] == 2
    assert summary["matched"] == 2

    rows = results_of(client, rubric_id)
    assert {item["student_no"] for item in rows} == {fixture.no_zhang, fixture.no_wang}
    assert all(item["source_filename"].endswith(".docx") for item in rows)

    # 限定班级只在该班匹配
    limited = results_of(client, rubric_id, class_id=fixture.class_soft)
    assert [item["student_no"] for item in limited] == [fixture.no_zhang]

    # 限定课程后，课程外的学生匹配不上
    scoped = upload(
        client,
        rubric_id,
        "补交.zip",
        zip_bytes({f"{fixture.department}-软件工程-软工2301-{fixture.no_li}-李四.docx": b"docx"}),
        course_id=fixture.course_id,
    )
    assert scoped["matched"] == 1
    assert scoped["created"] == 1


def test_upload_reports_skipped_files(client):
    setup(client)
    rubric_id = create_rubric(client, kind="homework", name="没有可评文件")
    summary = upload(
        client,
        rubric_id,
        "只有素材.zip",
        zip_bytes({"素材.png": b"\x89PNG", "说明.txt": "x"}),
    )
    assert summary["accepted_files"] == 0
    assert "没有找到可评审的文件" in summary["message"]
    assert len(summary["issues"]) == 2


def test_reupload_overwrites_and_resets(client):
    fixture = setup(client)
    rubric_id = create_rubric(client, kind="lab_report", name=f"实验二-{fixture.tag}")

    def payload(content: bytes) -> bytes:
        return zip_bytes(
            {f"张三_{fixture.no_zhang}.zip": zip_bytes({f"张三_{fixture.no_zhang}.pdf": content})}
        )

    assert upload(client, rubric_id, "批次.zip", payload(b"v1"))["created"] == 1
    grade_all(client, rubric_id)
    rows = results_of(client, rubric_id)
    assert rows[0]["status"] == "graded"
    assert rows[0]["score"] is not None

    second = upload(client, rubric_id, "批次2.zip", payload(b"v2-changed"))
    assert second["updated"] == 1
    assert second["created"] == 0

    rows = results_of(client, rubric_id)
    assert len(rows) == 1  # 还是一个学生一条记录
    assert rows[0]["status"] == "pending"  # 评审结果已重置
    assert rows[0]["score"] is None


# --------------------------------------------------------------------------- #
# 人工修正 / 删除 / 校验
# --------------------------------------------------------------------------- #
def test_manual_adjustment_and_delete(client):
    fixture = setup(client)
    rubric_id = create_rubric(client, kind="lab_report", name=f"实验三-{fixture.tag}")
    upload(
        client,
        rubric_id,
        "批次.zip",
        zip_bytes(
            {f"张三_{fixture.no_zhang}.zip": zip_bytes({f"张三_{fixture.no_zhang}.pdf": b"x"})}
        ),
    )
    grade_all(client, rubric_id)
    result = results_of(client, rubric_id)[0]

    response = client.patch(
        f"/api/grading-results/{result['id']}",
        json={"score": 95, "comment": "教师复核后调整"},
    )
    assert response.status_code == 200, response.text
    assert response.json()["score"] == 95
    assert response.json()["is_manual"] is True

    over = client.patch(f"/api/grading-results/{result['id']}", json={"score": 999}).json()
    assert over["score"] == 100  # 不能超过满分

    assert client.delete(f"/api/grading-results/{result['id']}").status_code == 204
    assert results_of(client, rubric_id) == []


def test_grading_requires_criteria(client):
    rubric_id = create_rubric(
        client, kind="homework", name="空细则", with_criteria=False, with_description=False
    )
    response = client.post(f"/api/rubrics/{rubric_id}/grade", json={})
    assert response.status_code == 400
    assert "评分标准" in response.json()["detail"]


def test_rubric_crud(client):
    rubric_id = create_rubric(client, kind="homework", name="待改名")
    updated = client.patch(
        f"/api/rubrics/{rubric_id}", json={"name": "改过名的作业", "total_score": 50}
    ).json()
    assert updated["name"] == "改过名的作业"
    assert updated["total_score"] == 50
    assert updated["stats"]["result_count"] == 0

    assert client.get("/api/rubrics", params={"kind": "homework"}).status_code == 200
    assert client.delete(f"/api/rubrics/{rubric_id}").status_code == 204
    assert client.get(f"/api/rubrics/{rubric_id}").status_code == 404


def test_dashboard_and_health(client):
    health = client.get("/api/health").json()
    assert health["status"] == "ok"
    assert health["storage"]["ok"] is True

    dashboard = client.get("/api/dashboard").json()
    assert dashboard["student_count"] >= 3
    assert dashboard["course_count"] >= 1
    assert "rubric_count" in dashboard
